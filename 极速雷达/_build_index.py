#!/usr/bin/env python3
"""极速雷达索引构建器 — 高性能版"""
import os, sys, json, time, zipfile
import xml.etree.ElementTree as ET
from concurrent.futures import ProcessPoolExecutor, as_completed

WORKSPACE = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INDEX     = os.path.join(WORKSPACE, ".workbuddy", "_radar_index.txt")
MANIFEST  = os.path.join(WORKSPACE, ".workbuddy", "_radar_manifest.json")
BIN_EXTS  = {'.pdf', '.docx', '.xlsx', '.xls'}
WORKERS   = 8

_W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
EXCLUDE_DIRS = {'node_modules', '.git', '.venv', 'venv', 'dist', 'build', 'target'}

def scan_pdf(fp):
    lines = []
    try:
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader
        
        reader = PdfReader(fp, strict=False)
        for i, p in enumerate(reader.pages, 1):
            t = p.extract_text()
            if not t: continue
            for j, ln in enumerate(t.split('\n')):
                ln = ln.strip()
                if ln:
                    lines.append(f"{fp}\tP{i}L{j+1}\t{ln}")
    except Exception:
        pass
    return lines

def scan_docx(fp):
    lines = []
    try:
        with zipfile.ZipFile(fp) as z:
            if 'word/document.xml' not in z.namelist(): return lines
            xml = z.read('word/document.xml')
        root = ET.fromstring(xml)
        for i, p_el in enumerate(root.iter(f'{{{_W_NS}}}p'), 1):
            text = "".join(t.text or "" for t in p_el.iter(f'{{{_W_NS}}}t')).strip()
            if text:
                lines.append(f"{fp}\tL{i}\t{text}")
    except Exception:
        pass
    return lines

def scan_xlsx(fp):
    lines = []
    try:
        import openpyxl
        wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
        for sn in wb.sheetnames:
            ws = wb[sn]
            for ri, row in enumerate(ws.iter_rows(values_only=True), 1):
                for ci, val in enumerate(row, 1):
                    if val is not None:
                        t = str(val).strip()
                        if t:
                            lines.append(f"{fp}\t{sn}:R{ri}C{ci}\t{t}")
    except Exception:
        try:
            with zipfile.ZipFile(fp) as z:
                if 'xl/sharedStrings.xml' not in z.namelist(): return lines
                xml = z.read('xl/sharedStrings.xml')
            _S_NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
            for i, si in enumerate(ET.fromstring(xml).findall(f'{{{_S_NS}}}si')):
                t = ''.join(t.text or '' for t in si.iter(f'{{{_S_NS}}}t')).strip()
                if t:
                    lines.append(f"{fp}\tR{i+1}\t{t}")
        except Exception:
            pass
    return lines

def scan_xls(fp):
    lines = []
    try:
        import xlrd
        wb = xlrd.open_workbook(fp, on_demand=True)
        for sn in wb.sheet_names():
            ws = wb[sn]
            for ri in range(ws.nrows):
                for ci in range(ws.ncols):
                    v = ws.cell_value(ri, ci)
                    if isinstance(v, str):
                        v = v.strip()
                    if v:
                        lines.append(f"{fp}\t{sn}:R{ri+1}C{ci+1}\t{str(v)}")
    except Exception:
        pass
    return lines

def scan_file(fp):
    ext = os.path.splitext(fp)[1].lower()
    if ext == '.pdf':   return scan_pdf(fp)
    elif ext == '.docx': return scan_docx(fp)
    elif ext == '.xlsx': return scan_xlsx(fp)
    elif ext == '.xls':  return scan_xls(fp)
    return []

def main():
    files = []
    manifest = {}
    for root, dirs, fns in os.walk(WORKSPACE):
        dirs[:] = [d for d in dirs if not d.startswith('.') and d not in EXCLUDE_DIRS]
        for fn in fns:
            if fn.startswith('~$'): continue
            fp = os.path.join(root, fn)
            ext = os.path.splitext(fn)[1].lower()
            if ext in BIN_EXTS:
                st = os.stat(fp)
                files.append(fp)
                manifest[fp] = [st.st_mtime, st.st_size]

    print(f"索引构建中: {len(files)} 个二进制文件...")
    t0 = time.time()
    all_lines = []

    with ProcessPoolExecutor(max_workers=WORKERS) as ex:
        futures = {ex.submit(scan_file, fp): fp for fp in files}
        for fut in as_completed(futures):
            all_lines.extend(fut.result())

    all_lines.sort()

    os.makedirs(os.path.dirname(INDEX), exist_ok=True)
    with open(INDEX, 'w', encoding='utf-8') as f:
        f.write(f"# 极速雷达索引 | {len(all_lines)} 行 | {time.strftime('%Y-%m-%d %H:%M:%S')}\n")
        f.write('\n'.join(all_lines))

    with open(MANIFEST, 'w') as f:
        json.dump(manifest, f)

    print(f"索引完成: {len(all_lines)} 行, 耗时 {time.time()-t0:.1f}s")

if __name__ == '__main__':
    main()
