#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
wide_table_query.py — 交易结算核心系统需求集合 大宽表查询
=====================================================
直接使用 HTTP API 接口快速获取数据，支持 Session 复用。
如果在 API 路径下鉴权失败或遇到异常，自动退回至 Playwright 浏览器模式（进行登录并提取最新 token，或执行传统物理导出）。

用法:
    .venv/bin/python wide_table_query.py <搜索关键字> [--headed] [--output result.md] [--html]
    .venv/bin/python wide_table_query.py PG202204-0155
    .venv/bin/python wide_table_query.py "融资融券"

依赖: playwright, openpyxl, requests
      虚拟环境 .venv/
"""

import sys
import os
import re
import time
import argparse
import json
import requests
from pathlib import Path
from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout
from openpyxl import load_workbook, Workbook

def decrypt_password(enc_str, key_path='~/.workbuddy/.meeting_skill_key'):
    """解密存储在代码中的加密密码"""
    if not enc_str or not isinstance(enc_str, str):
        return ""
    if not enc_str.startswith('ENC:'):
        return enc_str
    kp = os.path.expanduser(key_path)
    if not os.path.exists(kp):
        return enc_str
    try:
        from cryptography.fernet import Fernet
        with open(kp, 'rb') as f:
            key = f.read()
        fern = Fernet(key)
        return fern.decrypt(enc_str[4:].encode('utf-8')).decode('utf-8')
    except Exception:
        return enc_str

def load_credentials():
    username = os.environ.get("PLATFORM_USERNAME") or os.environ.get("FINTECH_USERNAME") or ""
    password = os.environ.get("PLATFORM_PASSWORD") or os.environ.get("FINTECH_PASSWORD") or ""
    platform_url = os.environ.get("PLATFORM_URL") or "https://fintech.gtht.com.cn"

    script_dir = os.path.dirname(os.path.abspath(__file__))
    skill_root = os.path.dirname(script_dir)

    config_paths = [
        os.path.join(os.getcwd(), "config.json"),
        os.path.join(skill_root, "config.json"),
        os.path.join(script_dir, "config.json"),
        "/Volumes/Macintosh HD_Data/WorkBuddy/需求管理/config.json",
        os.path.expanduser("~/.workbuddy/config.json"),
    ]

    for p in config_paths:
        if os.path.exists(p):
            try:
                with open(p, "r", encoding="utf-8") as f:
                    creds = json.load(f)
                    u = str(creds.get("username", "")).strip()
                    p_plain = str(creds.get("password", "")).strip()
                    p_enc = str(creds.get("password_encrypted", "")).strip()
                    url_val = str(creds.get("platform_url", "")).strip()

                    if url_val:
                        platform_url = url_val.rstrip("/")
                    if u and u not in ("YOUR_USERNAME", "你的工号") and not username:
                        username = u

                    if not password:
                        if p_plain and p_plain not in ("YOUR_PASSWORD", "YOUR_PLAIN_PASSWORD", "你的登录密码", "密码"):
                            password = p_plain
                        elif p_enc:
                            dec = decrypt_password(p_enc)
                            if dec:
                                password = dec
                if username and password:
                    break
            except Exception:
                pass

    if not username or not password:
        print("❌ [配置缺失错误] 未找到有效的科技平台登录凭据 (USERNAME / PASSWORD)", file=sys.stderr)

    return username, password, platform_url

USERNAME, PASSWORD, PLATFORM_URL = load_credentials()
LOGIN_URL = f"{PLATFORM_URL}/kjpt/user/login"

TABLE_URL = (
    f"{PLATFORM_URL}/kjpt/OnlineGrid"
    "?tableId=1597051499572236288"
    "&tableName=%E4%BA%A4%E6%98%93%E7%BB%93%E7%AE%97%E6%A0%B8%E5%BF%83%E7%B3%BB%E7%BB%9F%E9%9C%80%E6%B1%82%E9%9B%86%E5%90%88"
)

# Content API URL
CONTENT_URL = f"{PLATFORM_URL}/api/table-service/table/row/getSheetContent?sheetId=1597051509256884224"

# Session 文件 — 与其他 fintech 脚本共用（统一存放在工作目录）
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SESSION_FILE = "/Volumes/Macintosh HD_Data/WorkBuddy/需求管理/.fintech_session.json"

# 需要提取的关键列（中文名，用于 Excel 表头匹配）
KEY_COLUMNS = [
    "史诗编号&史诗名称",
    "需求编号",
    "需求名称",
    "Story系统",
    "Story状态",
    "Story业务验收结果",
    "是否加急",
    "计划生产排期",
    "版本排期",
    "备注",
]

# 列字段名映射（用于内部引用）
FIELD_KEYS = [
    "epicId",
    "demandId",
    "demandName",
    "storySystem",
    "testStatus",
    "storyResult",
    "isUrgent",
    "planProdDate",
    "versionDate",
    "remark",
]

# API 返回的字段与内部字段的映射
API_FIELD_MAP = {
    "epicConcat": "epicId",
    "demandId": "demandId",
    "demandName": "demandName",
    "storySystemName": "storySystem",
    "storyStatusName": "testStatus",
    "storyBusName": "storyResult",
    "demandUrgentStatus": "isUrgent",
    "planProdLineDate": "planProdDate",
    "storyPlanLineName": "versionDate",
    "beizhuText": "remark",
}

# ============================================================
# 工具函数
# ============================================================
def wait_stable(page, ms=1500):
    """等待页面稳定"""
    time.sleep(ms / 1000.0)


def _clean(s):
    """去除换行、多余空白"""
    if not s:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip()


def _display_width(s):
    """计算显示宽度：CJK/全角=2，ASCII/半角=1"""
    w = 0
    for ch in str(s):
        if ord(ch) >= 0x2E80:
            w += 2
        else:
            w += 1
    return w


def _pad_cjk(text, width):
    """左对齐填充到指定显示宽度"""
    cur = _display_width(text)
    if cur >= width:
        return str(text)
    return str(text) + " " * (width - cur)


def _truncate_cjk(text, width):
    """截断到指定显示宽度，末尾加 …"""
    text = str(text)
    if _display_width(text) <= width:
        return text
    result = []
    w = 0
    limit = width - 2
    for ch in text:
        ch_w = 2 if ord(ch) >= 0x2E80 else 1
        if w + ch_w > limit:
            break
        result.append(ch)
        w += ch_w
    return "".join(result) + "\u2026"


# ============================================================
# API 辅助逻辑
# ============================================================
def load_session_credentials():
    """从 session 文件提取 cookies 和 token"""
    cookies = {}
    token = None
    if not os.path.exists(SESSION_FILE):
        return cookies, token
    try:
        with open(SESSION_FILE, 'r', encoding='utf-8') as f:
            session_data = json.load(f)
        for c in session_data.get("cookies", []):
            if c.get("domain") in ("fintech.gtht.com.cn", ".gtht.com.cn"):
                cookies[c["name"]] = c["value"]
        for origin_data in session_data.get("origins", []):
            if origin_data.get("origin") == "https://fintech.gtht.com.cn":
                for item in origin_data.get("localStorage", []):
                    if item.get("name") == "GTJA_TOKEN":
                        token = item.get("value")
                        break
    except Exception as e:
        print(f"[API] 读取 session 文件失败: {e}")
    return cookies, token


def fetch_data_via_api(cookies, token, cache_file=None):
    """通过 HTTP API 直接获取表格全量数据"""
    headers = {
        "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        "Referer": "https://fintech.gtht.com.cn/kjpt/OnlineGrid?tableId=1597051499572236288",
        "accept": "application/json",
        "content-type": "application/json",
    }
    if token:
        headers["token"] = token
    
    r = requests.get(CONTENT_URL, cookies=cookies, headers=headers, timeout=60)
    if r.status_code != 200:
        raise ValueError(f"HTTP 状态码错误: {r.status_code}")
    
    resp_json = r.json()
    if resp_json.get("code") != 200:
        raise ValueError(f"API 返回错误: {resp_json.get('message') or resp_json.get('msg')}")
        
    # 如果指定了缓存路径，则异步/同步保存一份到本地
    if cache_file:
        try:
            os.makedirs(os.path.dirname(os.path.abspath(cache_file)), exist_ok=True)
            with open(cache_file, 'w', encoding='utf-8') as f:
                json.dump(resp_json, f, ensure_ascii=False)
        except Exception as ce:
            print(f"\n[CACHE] ⚠️ 保存本地缓存失败: {ce}")
        
    rows = resp_json.get("data", {}).get("data", [])
    return rows


def save_to_excel(rows, filepath):
    """将 API 结果保存为与浏览器导出格式兼容的轻量级 Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    
    # 写入表头
    ws.append(KEY_COLUMNS)
    
    # 写入行
    for r in rows:
        row_values = []
        for field_key in FIELD_KEYS:
            row_values.append(r.get(field_key, ""))
        ws.append(row_values)
        
    os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)
    wb.save(filepath)


# ============================================================
# 浏览器自动化与传统导出（作为兜底）
# ============================================================
_JS_LOGIN = "() => !!document.querySelector('input[placeholder*=\"工号\"]')"
_JS_FILTER_BAR = (
    "() => {"
    "  const b = document.querySelector('.ag-status-bar');"
    "  return b && /过滤后数据总数/.test(b.innerText || '');"
    "}"
)
_JS_ROWS_VISIBLE = (
    "() => {"
    "  const vp = document.querySelector('.ag-body-viewport') || document.querySelector('[class*=\"ag-root\"]');"
    "  if (!vp) return false;"
    "  return vp.querySelectorAll('.ag-row').length > 0;"
    "}"
)

def ensure_login(context, page):
    """确保登录态有效。已登录则跳过，未登录则执行登录并保存 session"""
    print("[登录]", end=" ", flush=True)
    t0 = time.time()

    try:
        page.goto(TABLE_URL, wait_until="domcontentloaded", timeout=20000)
        page.wait_for_timeout(1000)  # 给予页面1秒时间来处理重定向或载入前端路由
        # 如果当前URL不包含login，且localStorage已存在授权token，则代表已有登录态
        if "login" not in page.url and page.evaluate("() => !!localStorage.getItem('GTJA_TOKEN')"):
            print(f"已有登录态（{time.time() - t0:.1f}s）")
            return
    except Exception:
        pass

    # 需要登录
    print("登录中...", end=" ", flush=True)
    if "login" not in page.url:
        page.goto(LOGIN_URL, wait_until="domcontentloaded", timeout=30000)
    page.wait_for_selector('input[placeholder*="工号"]', timeout=15000)
    page.locator('input[placeholder*="工号"]').fill(USERNAME)
    page.locator('input[placeholder*="密码"]').fill(PASSWORD)
    page.locator('button:has-text("提 交")').click()
    
    # 等待重定向回到OnlineGrid页面并加载数据行（从而保证 localStorage 已被成功注入）
    try:
        page.wait_for_url("**/OnlineGrid**", timeout=30000)
        page.wait_for_function(_JS_ROWS_VISIBLE, polling=200, timeout=20000)
    except PWTimeout:
        pass
    
    time.sleep(1)

    try:
        context.storage_state(path=SESSION_FILE)
        print(f"完成（{time.time() - t0:.1f}s）")
    except Exception as e:
        print(f"完成（session 保存失败: {e}）")


def search_and_export(page, keyword, download_dir):
    """
    兜底的物理导出路径：
    1. 确保在表格页
    2. 在搜索框填入关键字，触发 AG Grid 过滤
    3. 点击导出 -> 导出excel -> 下载
    """
    print(f"[搜索] 关键字: {keyword} (使用浏览器物理导出中...)")
    t0 = time.time()

    # 确保在表格页
    if "OnlineGrid" not in page.url:
        page.goto(TABLE_URL, wait_until="domcontentloaded", timeout=60000)

    # 等待 AG Grid 渲染就绪
    try:
        page.wait_for_function(_JS_ROWS_VISIBLE, polling=200, timeout=20000)
    except PWTimeout:
        print("[搜索] ⚠️ 等待数据行超时，继续...")

    print(f"[搜索] 表格就绪（{time.time() - t0:.1f}s）")

    # 填入搜索框
    try:
        search_box = page.locator('#filter-text-box')
        search_box.wait_for(state="visible", timeout=10000)

        # 瞬间填充并触发事件
        search_box.click()
        search_box.fill(keyword)
        print("[搜索] 关键字已填入，等待过滤生效...")
        
        # 等待过滤完成
        page.wait_for_function(_JS_FILTER_BAR, polling=200, timeout=15000)

        # 验证过滤结果
        grid_info = page.evaluate("""() => {
            const grid = document.querySelector('.ag-body-viewport') ||
                         document.querySelector('[class*="ag-root"]');
            const rows = grid ? grid.querySelectorAll('.ag-row') : [];
            const statusBar = document.querySelector('.ag-status-bar') ||
                             document.querySelector('[ref="statusBar"]');
            let statusText = '';
            if (statusBar) statusText = statusBar.innerText;
            return { visibleRows: rows.length, statusText: statusText };
        }""")
        print(f"[搜索] 过滤完成 — 可见行: {grid_info['visibleRows']}, 状态: {grid_info['statusText'].strip()} ({time.time() - t0:.1f}s)")

    except PWTimeout:
        print("[搜索] 搜索框或过滤超时，继续尝试直接导出")
    except Exception as e:
        print(f"[搜索] 搜索过程异常: {e}，继续导出")

    # 两步导出：主按钮 → popover → "导出excel"
    print("[导出] 打开导出菜单...")
    os.makedirs(download_dir, exist_ok=True)

    export_btn = page.locator('button:has-text("导出")').first
    export_btn.click()
    
    # 等待 popover 菜单中的 "导出excel" 按钮可见
    page.wait_for_selector('button:has-text("导出excel")', state="visible", timeout=5000)

    # 点击 popover 内的 "导出excel" 并监听下载
    print("[导出] 点击「导出excel」...")
    excel_btn = page.locator('button:has-text("导出excel")').first
    with page.expect_download(timeout=120000) as download_info:
        excel_btn.click()

    download = download_info.value
    save_path = os.path.join(download_dir, f"exported_{keyword or 'all'}.xlsx")
    download.save_as(save_path)

    file_size = os.path.getsize(save_path)
    print(f"[导出] 文件已保存: {save_path} ({file_size / 1024 / 1024:.2f} MB)")
    return save_path


def read_and_filter_excel(filepath, keyword=None):
    """从本地 Excel 文件中二次读取并匹配数据行"""
    print(f"[读取] 加载 Excel: {filepath}")
    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active

    # 读取表头
    header_row = [_clean(cell.value) for cell in ws[1]]

    # 建立关键列索引映射
    col_map = {}
    missing_cols = []
    for col_name, field_key in zip(KEY_COLUMNS, FIELD_KEYS):
        if col_name in header_row:
            col_map[field_key] = header_row.index(col_name)
        else:
            missing_cols.append(col_name)

    if not col_map:
        raise ValueError("未找到匹配的关键列，请检查 Excel 表头")

    # 读取数据行并过滤
    matched = []
    total = 0
    kw_lower = keyword.lower() if keyword else None

    for row in ws.iter_rows(min_row=2, values_only=True):
        total += 1
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue

        row_data = {}
        for field_key, col_idx in col_map.items():
            row_data[field_key] = _clean(row[col_idx]) if col_idx < len(row) else ""

        if kw_lower:
            haystack = " ".join(row_data.values()).lower()
            if kw_lower not in haystack:
                continue

        matched.append(row_data)

    wb.close()
    print(f"[读取] Excel 共 {total} 行，有效匹配 {len(matched)} 行")
    return matched


# ============================================================
# 格式化输出
# ============================================================
def format_compact(rows, keyword):
    """对话窗口精简输出：列对齐"""
    if not rows:
        return f"**交易结算核心系统需求集合　|　搜索：{keyword}**\n\n无匹配结果。\n"

    # 按 需求编号+Story系统 去重
    seen = set()
    unique = []
    for r in rows:
        key = (r.get("demandId", ""), r.get("storySystem", ""))
        if key not in seen:
            seen.add(key)
            unique.append(r)

    # 按需求编号分组
    by_demand = {}
    for r in unique:
        did = r.get("demandId", "")
        if did not in by_demand:
            by_demand[did] = []
        by_demand[did].append(r)

    lines = []
    lines.append(f"**交易结算核心系统需求集合　|　搜索：{keyword}**")
    lines.append(f"需求数：{len(by_demand)}　|　记录数：{len(unique)}")

    # 表头
    cols = ["序号", "需求编号", "需求名称", "Story系统", "Story状态", "验收结果", "生产排期", "版本排期"]
    col_keys = ["idx", "did", "name", "system", "test", "verify", "plan", "version"]
    col_w = {k: _display_width(h) for k, h in zip(col_keys, cols)}

    # 收集数据 + 计算最大宽度
    data_rows = []
    seq = 0
    for did, group in by_demand.items():
        seq += 1
        first = group[0]
        data_rows.append({
            "idx": f"{seq}.",
            "did": _clean(did),
            "name": _clean(first.get("demandName", "")),
            "system": _clean(first.get("storySystem", "")),
            "test": _clean(first.get("testStatus", "")),
            "verify": _clean(first.get("storyResult", "")),
            "plan": _clean(first.get("planProdDate", "")),
            "version": _clean(first.get("versionDate", "")),
        })
        for extra in group[1:]:
            seq += 1
            data_rows.append({
                "idx": f"{seq}.",
                "did": "",
                "name": "",
                "system": _clean(extra.get("storySystem", "")),
                "test": _clean(extra.get("testStatus", "")),
                "verify": _clean(extra.get("storyResult", "")),
                "plan": _clean(extra.get("planProdDate", "")),
                "version": _clean(extra.get("versionDate", "")),
            })

    for k in col_keys:
        for dr in data_rows:
            col_w[k] = max(col_w[k], _display_width(dr.get(k, "")))

    SEP = "  "
    lines.append(SEP.join(_pad_cjk(h, col_w[k]) for h, k in zip(cols, col_keys)))

    many = len(data_rows) > 30
    for dr in data_rows:
        parts = []
        for k in col_keys:
            v = dr.get(k, "")
            if k == "name" and many:
                parts.append(_truncate_cjk(v, col_w[k]))
            else:
                parts.append(_pad_cjk(v, col_w[k]))
        lines.append(SEP.join(parts))

    # 备注
    remarks = [
        (r.get("demandId", ""), _clean(r.get("remark", "")))
        for r in unique if _clean(r.get("remark", ""))
    ]
    if remarks:
        lines.append("")
        lines.append("**备注：**")
        for did, rm in remarks:
            if did:
                lines.append(f"　　{did}：{rm}")
            else:
                lines.append(f"　　{rm}")

    lines.append(f"共 {len(by_demand)} 个需求、{len(unique)} 条记录")
    return "\n".join(lines)


def format_markdown(rows, keyword):
    """完整 Markdown 输出（文件用）"""
    if not rows:
        return f"**交易结算核心系统需求集合　|　搜索：{keyword}**\n\n无匹配结果。\n"

    # 去重
    seen = set()
    unique = []
    for r in rows:
        key = (r.get("demandId", ""), r.get("storySystem", ""))
        if key not in seen:
            seen.add(key)
            unique.append(r)

    # 按需求编号分组
    by_demand = {}
    for r in unique:
        did = r.get("demandId", "")
        if did not in by_demand:
            by_demand[did] = []
        by_demand[did].append(r)

    lines = []
    lines.append(f"**交易结算核心系统需求集合　|　搜索：{keyword}**")
    lines.append(f"**需求数：{len(by_demand)}　|　记录数：{len(unique)}**")
    lines.append("")

    detail_url_tmpl = "https://fintech.gtht.com.cn/kjpt/DemandManage/details?demandId={}"

    for i, (did, group) in enumerate(by_demand.items(), 1):
        first = group[0]
        link = detail_url_tmpl.format(did)
        name = _clean(first.get("demandName", ""))
        epic = _clean(first.get("epicId", ""))
        title_extra = f"  {name}" if name else ""

        lines.append(f"**{i}. [{did}]({link})**{title_extra}")
        if epic:
            lines.append(f"　　史诗：{epic}")

        for j, r in enumerate(group):
            sys_name = _clean(r.get("storySystem", ""))
            test = _clean(r.get("testStatus", ""))
            verify = _clean(r.get("storyResult", ""))
            urgent = _clean(r.get("isUrgent", ""))
            plan = _clean(r.get("planProdDate", ""))
            ver = _clean(r.get("versionDate", ""))
            rm = _clean(r.get("remark", ""))

            lines.append(f"　　{sys_name}")

            meta = []
            if test:   meta.append(f"交易组测试状态：{test}")
            if verify: meta.append(f"Story业务验收结果：{verify}")
            if urgent: meta.append(f"是否加急：{urgent}")
            if plan:   meta.append(f"计划生产排期：{plan}")
            if ver:    meta.append(f"版本排期：{ver}")
            if rm:     meta.append(f"备注：{rm}")
            if meta:
                lines.append(f"　　　{'　|　'.join(meta)}")

        lines.append("")

    lines.append("---")
    lines.append(f"**共 {len(by_demand)} 个需求、{len(unique)} 条记录，搜索关键字：{keyword}**")
    return "\n".join(lines)


def _status_color(test):
    """根据测试状态返回 CSS class 名"""
    if "通过" in test or test == "结束":
        return "pass"
    if "自测" in test:
        return "testing"
    if "排期" in test or "受理" in test:
        return "pending"
    if "终止" in test:
        return "fail"
    return "default"


def _verify_color(verify):
    """根据验收结果返回 CSS class 名"""
    if "通过" in verify:
        return "pass"
    if "无需" in verify:
        return "skip"
    if "待" in verify:
        return "pending"
    return "default"


def format_html(rows, keyword):
    """生成带卡片可视化的 HTML 文件"""
    if not rows:
        return f'<div style="color:var(--color-text-secondary);font-size:13px;">无匹配结果。</div>'

    seen = set()
    unique = []
    for r in rows:
        key = (r.get("demandId", ""), r.get("storySystem", ""))
        if key not in seen:
            seen.add(key)
            unique.append(r)

    by_demand = {}
    for r in unique:
        did = r.get("demandId", "")
        if did not in by_demand:
            by_demand[did] = []
        by_demand[did].append(r)

    detail_url_tmpl = "https://fintech.gtht.com.cn/kjpt/DemandManage/details?demandId={}"

    # 统计
    total_systems = len(unique)
    verify_pass = sum(1 for r in unique if "通过" in r.get("storyResult", ""))
    verify_total = sum(1 for r in unique if r.get("storyResult", ""))
    plan_dates = set(r.get("planProdDate", "") for r in unique if r.get("planProdDate", ""))

    epic = _clean(unique[0].get("epicId", "")) if unique else ""
    epic_short = epic.split("+")[0] if epic else keyword
    epic_name = epic.split("+")[1] if "+" in epic else ""

    # 计划排期显示
    plan_display = ""
    if len(plan_dates) == 1:
        d = plan_dates.pop()
        plan_display = d.replace("2026-", "") if d.startswith("2026-") else d
    elif plan_dates:
        plan_display = "多排期"

    parts = []

    # CSS
    parts.append("""<style>
*{margin:0;padding:0;box-sizing:border-box}
body{font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Helvetica,Arial,sans-serif;
  color:#2C2C2A;background:#F1EFE8;line-height:1.6;padding:24px;max-width:720px;margin:0 auto}
.card{background:#fff;border-radius:10px;border:0.5px solid rgba(0,0,0,0.12);padding:14px 16px;margin-bottom:10px}
.metric{background:#D3D1C7;border-radius:6px;padding:10px 12px;text-align:center}
.metric .label{font-size:11px;color:#888780}
.metric .value{font-size:22px;font-weight:500;color:#2C2C2A;margin-top:2px}
.metric .value.green{color:#3B6D11}
.metrics{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;margin:16px 0 20px}
.header{display:flex;align-items:center;gap:10px;margin-bottom:4px}
.header .icon{width:36px;height:36px;border-radius:50%;background:#B5D4F4;display:flex;align-items:center;justify-content:center;flex-shrink:0}
.header .icon svg{width:18px;height:18px;stroke:#0C447C;fill:none;stroke-width:2;stroke-linecap:round;stroke-linejoin:round}
.header .title{font-size:15px;font-weight:500;color:#2C2C2A}
.header .sub{font-size:12px;color:#888780;margin-top:1px}
.req-title{font-size:13px;font-weight:500;color:#2C2C2A;flex:1}
.req-link{font-size:11px;color:#185FA5;white-space:nowrap;text-decoration:none;margin-left:10px}
.req-head{display:flex;align-items:flex-start;margin-bottom:8px}
.sys-row{display:flex;align-items:center;gap:8px;flex-wrap:wrap;margin-bottom:4px}
.sys-tag{font-size:11px;padding:2px 8px;border-radius:4px;background:#B5D4F4;color:#0C447C}
.status-text{font-size:11px;color:#888780}
.badge{font-size:11px;padding:1px 6px;border-radius:3px}
.badge.pass{background:#EAF3DE;color:#3B6D11}
.badge.testing{background:#FAEEDA;color:#854F0B}
.badge.pending{background:#E1F5EE;color:#0F6E56}
.badge.fail{background:#FCEBEB;color:#A32D2D}
.badge.skip{background:#F1EFE8;color:#888780}
.badge.default{background:#F1EFE8;color:#888780}
.warn{margin-top:14px;padding:10px 14px;background:#FAEEDA;border-radius:6px;display:flex;align-items:center;gap:8px}
.warn svg{width:14px;height:14px;stroke:#854F0B;fill:none;stroke-width:2;flex-shrink:0}
.warn p{font-size:12px;color:#854F0B}
.footer{text-align:center;font-size:11px;color:#888780;margin-top:16px}
</style>""")

    # Header
    parts.append("""<div class="header">
  <div class="icon"><svg viewBox="0 0 24 24"><rect x="3" y="3" width="7" height="7"/><rect x="14" y="3" width="7" height="7"/><rect x="3" y="14" width="7" height="7"/><rect x="14" y="14" width="7" height="7"/></svg></div>
  <div><div class="title">交易结算核心系统需求集合</div><div class="sub">""")
    parts.append(f"{epic_short}")
    if epic_name:
        parts.append(f"+{epic_name}")
    parts.append(f"</div></div></div>")

    # Metrics
    parts.append(f"""<div class="metrics">
  <div class="metric"><div class="label">需求</div><div class="value">{len(by_demand)}</div></div>
  <div class="metric"><div class="label">系统记录</div><div class="value">{total_systems}</div></div>
  <div class="metric"><div class="label">验收通过</div><div class="value green">{verify_pass}/{verify_total}</div></div>
  <div class="metric"><div class="label">排期</div><div class="value">{plan_display}</div></div>
</div>""")

    # Demand cards
    for i, (did, group) in enumerate(by_demand.items(), 1):
        first = group[0]
        name = _clean(first.get("demandName", ""))
        link = detail_url_tmpl.format(did)
        if _display_width(name) > 40:
            name = _truncate_cjk(name, 40)

        parts.append(f'<div class="card">')
        parts.append(f'<div class="req-head"><div class="req-title">{name}</div><a class="req-link" href="{link}" target="_blank">{did} &#8599;</a></div>')

        for r in group:
            sys_name = _clean(r.get("storySystem", ""))
            test = _clean(r.get("testStatus", ""))
            verify = _clean(r.get("storyResult", ""))

            parts.append(f'<div class="sys-row">')
            parts.append(f'<span class="sys-tag">{sys_name}</span>')
            if test:
                tc = _status_color(test)
                if tc in ("pass", "fail"):
                    parts.append(f'<span class="badge {tc}">{test}</span>')
                else:
                    parts.append(f'<span class="status-text">{test}</span>')
            if verify:
                vc = _verify_color(verify)
                parts.append(f'<span class="badge {vc}">{verify}</span>')
            parts.append(f'</div>')

        parts.append(f'</div>')

    # Warning remarks
    remarks = []
    seen_rm = set()
    for r in unique:
        rm = _clean(r.get("remark", ""))
        if rm and rm not in seen_rm:
            seen_rm.add(rm)
            remarks.append(rm)
    if remarks:
        rm_text = "；".join(remarks)
        parts.append(f"""<div class="warn">
  <svg viewBox="0 0 24 24"><circle cx="12" cy="12" r="10"/><line x1="12" y1="8" x2="12" y2="12"/><line x1="12" y1="16" x2="12.01" y2="16"/></svg>
  <p>{rm_text}</p>
</div>""")

    parts.append(f'<div class="footer">共 {len(by_demand)} 个需求、{total_systems} 条记录 | 搜索：{keyword}</div>')

    return "\n".join(parts)


# ============================================================
# 主函数
# ============================================================
def main():
    parser = argparse.ArgumentParser(description="交易结算核心系统需求集合 — 大宽表查询")
    parser.add_argument("keyword", help="搜索关键字（需求编号、需求名称、系统名等）")
    parser.add_argument("--headed", action="store_true", help="显示浏览器窗口")
    parser.add_argument("--output", "-o", help="输出 Markdown 文件路径")
    parser.add_argument("--html", action="store_true", help="同时生成 HTML 可视化 file")
    parser.add_argument("--fresh", action="store_true", help="强制从服务器获取最新数据，不使用本地缓存")
    args = parser.parse_args()

    keyword = args.keyword.strip()
    if not keyword:
        print("错误：请输入搜索关键字")
        sys.exit(1)

    print(f"\n{'='*50}")
    print(f"  交易结算核心系统需求集合 — 大宽表查询 (高速 API 版)")
    print(f"  搜索: {keyword}")
    print(f"{'='*50}")

    # 下载目录
    download_dir = os.path.join(SCRIPT_DIR, "downloads")
    excel_path = os.path.join(download_dir, f"exported_{keyword or 'all'}.xlsx")
    
    rows = []
    query_success = False

    t_start = time.time()

    # ============================================================
    # Cache Check: 本地缓存校验 (缓存5分钟)
    # ============================================================
    CACHE_FILE = os.path.join(download_dir, ".sheet_cache.json")
    cache_duration = 300  # 5分钟

    if not args.fresh and os.path.exists(CACHE_FILE):
        cache_age = time.time() - os.path.getmtime(CACHE_FILE)
        if cache_age < cache_duration:
            try:
                print(f"[CACHE] 正在从本地缓存读取数据 (缓存年龄: {cache_age:.0f}s)...", end=" ", flush=True)
                t_cache_start = time.time()
                with open(CACHE_FILE, 'r', encoding='utf-8') as f:
                    cache_json = json.load(f)
                raw_rows = cache_json.get("data", {}).get("data", [])
                
                # 过滤匹配
                kw_lower = keyword.lower()
                for r in raw_rows:
                    row_data = {}
                    for api_key, internal_key in API_FIELD_MAP.items():
                        row_data[internal_key] = _clean(r.get(api_key, ""))
                    
                    haystack = " ".join(row_data.values()).lower()
                    if kw_lower in haystack:
                        rows.append(row_data)
                
                # 保存为兼容的本地 Excel 备份
                save_to_excel(rows, excel_path)
                query_success = True
                print(f"成功！用时: {time.time() - t_cache_start:.2f}s (使用 --fresh 可强制刷新)")
            except Exception as e:
                print(f"失败 ({e})，将重新从接口获取...")
                rows = []

    # ============================================================
    # Layer 1: 直接使用已有 Session Cookies + Token 进行 HTTP API 请求
    # ============================================================
    if not query_success:
        cookies, token = load_session_credentials()
        if cookies and token:
            try:
                print("[API] 正在尝试直接通过接口高速获取数据...", end=" ", flush=True)
                raw_rows = fetch_data_via_api(cookies, token, cache_file=CACHE_FILE)
                
                # 本地内存过滤
                kw_lower = keyword.lower()
                for r in raw_rows:
                    row_data = {}
                    for api_key, internal_key in API_FIELD_MAP.items():
                        row_data[internal_key] = _clean(r.get(api_key, ""))
                    
                    haystack = " ".join(row_data.values()).lower()
                    if kw_lower in haystack:
                        rows.append(row_data)
                
                # 保存本地兼容性 Excel 备份
                save_to_excel(rows, excel_path)
                query_success = True
                print(f"成功！用时: {time.time() - t_start:.2f}s (Layer 1)")
                
            except Exception as e:
                print(f"失败 ({e})")
                rows = []

    # ============================================================
    # Layer 2 & 3: Playwright 浏览器兜底
    # ============================================================
    if not query_success:
        print("[API] 缓存失效或直接查询失败，启动浏览器进行获取...")
        t_browser_start = time.time()
        
        with sync_playwright() as p:
            browser = p.chromium.launch(
                headless=not args.headed,
                args=["--disable-blink-features=AutomationControlled"]
            )
            context = browser.new_context(
                storage_state=SESSION_FILE if os.path.exists(SESSION_FILE) else None,
                viewport={"width": 1440, "height": 900},
                locale="zh-CN",
                accept_downloads=True,
            )
            page = context.new_page()

            try:
                # 确保完成登录并提取新 Session
                ensure_login(context, page)
                wait_stable(page, 3000)

                # 提取新 session token
                cookies_list = context.cookies()
                cookies = {c["name"]: c["value"] for c in cookies_list if c.get("domain") in ("fintech.gtht.com.cn", ".gtht.com.cn")}
                
                token = None
                try:
                    token = page.evaluate("() => localStorage.getItem('GTJA_TOKEN')")
                except Exception:
                    pass

                # Layer 2: 使用新登录的 Token 发送接口请求
                if cookies and token:
                    try:
                        print("[API] 登录成功，正在尝试通过新 Token 接口获取数据...", end=" ", flush=True)
                        t_layer2_start = time.time()
                        raw_rows = fetch_data_via_api(cookies, token, cache_file=CACHE_FILE)
                        
                        kw_lower = keyword.lower()
                        for r in raw_rows:
                            row_data = {}
                            for api_key, internal_key in API_FIELD_MAP.items():
                                row_data[internal_key] = _clean(r.get(api_key, ""))
                            
                            haystack = " ".join(row_data.values()).lower()
                            if kw_lower in haystack:
                                rows.append(row_data)
                                
                        save_to_excel(rows, excel_path)
                        query_success = True
                        print(f"成功！用时: {time.time() - t_layer2_start:.2f}s (Layer 2)")
                    except Exception as api_err:
                        print(f"接口尝试失败 ({api_err})，将启动物理导出...")
                        rows = []

                # Layer 3: 传统物理导出 Excel 方式（终极兜底）
                if not query_success:
                    excel_path = search_and_export(page, keyword, download_dir)
                    rows = read_and_filter_excel(excel_path, keyword)
                    query_success = True
                    print(f"[浏览器] 已完成物理导出并加载完毕 (Layer 3)")

            except Exception as e:
                print(f"\n[浏览器] 自动化运行遇到致命错误: {e}")
                import traceback
                traceback.print_exc()
                sys.exit(1)
            finally:
                try:
                    context.storage_state(path=SESSION_FILE)
                except Exception:
                    pass
                browser.close()

    if not query_success or not os.path.exists(excel_path):
        print("\n错误：大宽表数据查询彻底失败，未获取到任何结果")
        sys.exit(1)

    # 格式化输出
    output_compact = format_compact(rows, keyword)
    output_md = format_markdown(rows, keyword)

    # 对话窗口输出
    print("\n" + "=" * 60)
    print(output_compact)
    print("=" * 60)

    # 写输出文件
    if args.output:
        os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(output_md)
            f.write("\n\n---\n\n")
            f.write("## 对话窗口显示\n\n")
            f.write(output_compact)
            f.write("\n")
        print(f"\n结果已写入: {args.output}")

    # 默认也写一份 Markdown 结果
    default_output = os.path.join(
        download_dir, f"result_{keyword}.md"
    )
    with open(default_output, "w", encoding="utf-8") as f:
        f.write(output_md)
        f.write("\n\n---\n\n")
        f.write("## 对话窗口显示\n\n")
        f.write(output_compact)
        f.write("\n")
    print(f"\n结果已写入: {default_output}")

    # HTML 可视化
    if args.html:
        output_html = format_html(rows, keyword)
        html_path = os.path.join(download_dir, f"result_{keyword}.html")
        with open(html_path, "w", encoding="utf-8") as f:
            f.write("<!DOCTYPE html><html lang=\"zh-CN\"><head><meta charset=\"UTF-8\">")
            f.write("<meta name=\"viewport\" content=\"width=device-width,initial-scale=1\">")
            f.write(f"<title>需求集合 - {keyword}</title></head><body>")
            f.write(output_html)
            f.write("</body></html>")
        print(f"HTML 可视化: {html_path}")

    print(f"\n[完成] 任务全部结束，总执行耗时: {time.time() - t_start:.2f}秒\n")


if __name__ == "__main__":
    main()
