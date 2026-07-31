#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
国泰海通科技平台 - 集中交易平台持续建设项目
自动化 SPB 导出工具
============================================================

流水线：
  Phase 1  数据抓取 — Playwright 自动登录 fintech 平台，导出原始 Excel
  Phase 2  数据筛选 — 按「计划生产排期」+「版本排期」双条件过滤
  Phase 3  文件生成 — 主文件 / 增量报告 / 自动备份
  Phase 4  模板处理 — CLI / DLL / Table / Proc / HisTable / HisRunTable /
                     TableDevelop / UpgradeRemark / Remark / Init

用法：
  python auto_spb_export.py <版本号>
  示例：python auto_spb_export.py 20260522

依赖安装：
  pip install playwright pandas pypinyin openpyxl
  playwright install chromium
"""

import os
import sys
import re
import time
import shutil
from datetime import datetime

import openpyxl
import pandas as pd
from pypinyin import pinyin, Style
from playwright.sync_api import sync_playwright


def decrypt_password(enc_str, key_path='~/.workbuddy/.meeting_skill_key'):
    """解密存储在代码中的加密密码"""
    if not enc_str.startswith('ENC:'):
        return enc_str
    kp = os.path.expanduser(key_path)
    if not os.path.exists(kp):
        return enc_str
    try:
        from cryptography.fernet import Fernet
        with open(kp, 'rb') as f:
            key = f.read()
        fern = Fernet(key)
        return fern.decrypt(enc_str[4:].encode('utf-8')).decode('utf-8')
    except Exception:
        return enc_str


# ============================================================
#  配置区域
# ============================================================
BASE_VERSION = "SPB_V2.2.19"

USERNAME = "125360"
PASSWORD = decrypt_password("ENC:gAAAAABqS3nG3HaebAdXuA4oG2aGH1N93TKWM0WvB9czGsKVsKccervTuVmO8qIYLf-yuYBL2X-El9OOHp4W7HhUY3Yeg39rkg==")

TARGET_URL = (
    "https://fintech.gtht.com.cn/kjpt/OnlineGrid"
    "?tableId=1588040959688577024"
    "&tableName=%E9%9B%86%E4%B8%AD%E4%BA%A4%E6%98%93%E5%B9%B3%E5%8F%B0%E6%8C%81%E7%BB%AD%E5%BB%BA%E8%AE%BE%E9%A1%B9%E7%9B%AE"
)

LOGIN_URL = "https://fintech.gtht.com.cn/kjpt/user/login"

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
EXPORT_DIR = SCRIPT_DIR
DOWNLOAD_DIR = os.path.join(SCRIPT_DIR, "downloads")
OUTPUT_TXT = "excel_output.txt"
OUTPUT_UPT_TXT = "upt_excel_output.txt"
AUTH_STATE_PATH = os.path.join(SCRIPT_DIR, "auth_state.json")
# ============================================================


# ============================================================
#  工具函数
# ============================================================
def ensure_dir(directory):
    if not os.path.exists(directory):
        os.makedirs(directory)
        print(f"[INFO] 创建目录: {directory}")


def parse_version():
    """从命令行参数解析版本号和排期

    支持模式：
      python auto_spb_export.py 20260529                       # 单参数：计划生产排期=版本排期
      python auto_spb_export.py 20260529 20260529hot           # 双参数：分别指定
      python auto_spb_export.py 20260529hot --ver-only         # 仅按版本排期筛选（忽略计划生产排期）
    """
    if len(sys.argv) < 2:
        print("[ERROR] 缺少参数")
        print("  用法: python auto_spb_export.py <计划生产排期> [版本排期] [--ver-only]")
        print("  1 参数: python auto_spb_export.py 20260529              → 两者相同")
        print("  2 参数: python auto_spb_export.py 20260529 20260529hot   → 分别指定")
        print("  仅版本: python auto_spb_export.py 20260529hot --ver-only → 只按版本排期筛选")
        sys.exit(1)

    ver_only = '--ver-only' in sys.argv
    args = [a for a in sys.argv[1:] if a != '--ver-only']

    if len(args) == 1:
        if ver_only:
            # 仅版本排期模式
            ver_date = args[0].strip()
            plan_date = None
            version = ver_date
            print(f"[INFO] 仅版本排期模式: 版本排期={ver_date}, 版本号={version} (忽略计划生产排期)")
        else:
            version = args[0].strip()
            if not version.isdigit() or len(version) != 8:
                print(f"[ERROR] 版本号格式错误: {version}，应为 8 位数字，如 20260529")
                sys.exit(1)
            plan_date = version
            ver_date = version
    elif len(args) >= 2:
        plan_date = args[0].strip()
        ver_date = args[1].strip()
        if ver_date.lower() == "upt" or ver_date.lower().endswith("_upt"):
            ver_date = plan_date
        version = ver_date
        if not ver_only:
            print(f"[INFO] 双参数模式: 计划生产排期={plan_date}, 版本排期={ver_date}, 版本号={version}")
        else:
            plan_date = None
            print(f"[INFO] 仅版本排期模式: 版本排期={ver_date}, 版本号={version} (忽略计划生产排期)")
    return version, plan_date, ver_date, ver_only


# === 拼音工具 ===
def name_to_pinyin(name_str):
    if pd.isna(name_str) or not str(name_str).strip():
        return ""
    name_str = str(name_str).strip()
    if '-' not in name_str:
        return name_str.lower()
    chinese_name = name_str.split('-')[0]
    try:
        py_list = pinyin(chinese_name, style=Style.NORMAL)
        return ''.join([item[0] for item in py_list]).lower()
    except Exception:
        return chinese_name.lower()


def parse_path_string(path_str):
    """
    路径解析：逗号/分号/顿号/空格 → [['dir','file'],...]
    """
    if not path_str or str(path_str).lower() == 'nan':
        return None
    items = re.split(r'[;,；，、\s]+', str(path_str))
    path_items = []
    for item in items:
        item = item.strip()
        if not item:
            continue
        if '/' in item:
            parts = item.split('/', 1)
            d, f = parts[0].strip(), parts[1].strip() if len(parts) > 1 else ''
            path_items.append(f"['{d}','{f}']")
        else:
            path_items.append(f"['{item}','']")
    if not path_items:
        return None
    return "[" + ",".join(path_items) + "]"


# ============================================================
#  Phase 1 — 数据抓取
# ============================================================
def login_and_save(page, context):
    print("进行登录...")
    page.goto(LOGIN_URL, timeout=60000)
    page.fill('#workno', USERNAME)
    page.fill('#password', PASSWORD)
    page.click('button.ant-btn-primary')
    page.wait_for_load_state('networkidle')
    context.storage_state(path=AUTH_STATE_PATH)
    print("[OK] 登录成功并已保存会话状态")


def export_from_platform(version):
    print("\n" + "=" * 60)
    print("Phase 1: 从平台导出数据")
    print("=" * 60)

    ensure_dir(DOWNLOAD_DIR)

    with sync_playwright() as p:
        print("启动浏览器...")
        try:
            browser = p.chromium.launch(channel="chrome", headless=True)
        except Exception:
            # 兼容未安装 Chrome 但有 Chromium 的环境
            print("[INFO] 无法以 Chrome 运行，正在使用默认 Chromium 启动...")
            browser = p.chromium.launch(headless=True)

        # 尝试加载缓存的登录状态
        has_auth = os.path.exists(AUTH_STATE_PATH)
        if has_auth:
            print("发现历史登录缓存，尝试免密登录...")
            context = browser.new_context(
                storage_state=AUTH_STATE_PATH,
                viewport={'width': 1920, 'height': 1080},
                ignore_https_errors=True,
                accept_downloads=True,
            )
        else:
            print("未发现登录缓存...")
            context = browser.new_context(
                viewport={'width': 1920, 'height': 1080},
                ignore_https_errors=True,
                accept_downloads=True,
            )

        page = context.new_page()

        # 优化：屏蔽图片、媒体、字体请求以加速加载
        def intercept_route(route):
            if route.request.resource_type in ["image", "media", "font"]:
                route.abort()
            else:
                route.continue_()
        page.route("**/*", intercept_route)

        try:
            if not has_auth:
                login_and_save(page, context)
                print("访问表格页面...")
                page.goto(TARGET_URL, timeout=60000)
                page.wait_for_load_state('networkidle')
            else:
                try:
                    print("直接跳转目标表格页...")
                    page.goto(TARGET_URL, timeout=30000)
                    page.wait_for_load_state('networkidle')
                    if "login" in page.url:
                        print("[WARN] 登录状态失效，重新登录中...")
                        login_and_save(page, context)
                        print("访问表格页面...")
                        page.goto(TARGET_URL, timeout=60000)
                        page.wait_for_load_state('networkidle')
                except Exception as e:
                    print(f"[WARN] 快速访问失败 ({e})，重新执行完整登录...")
                    login_and_save(page, context)
                    print("访问表格页面...")
                    page.goto(TARGET_URL, timeout=60000)
                    page.wait_for_load_state('networkidle')

            page.wait_for_selector('text=一键展开', timeout=30000)

            # 尝试清除可能残留的过滤条件
            clear_btn = page.query_selector('text=清除所有筛选')
            if clear_btn:
                print("清除可能存在的历史过滤筛选条件...")
                clear_btn.click()
                time.sleep(1)
                confirm_btn = page.query_selector('button:has-text("确 定")')
                if confirm_btn:
                    confirm_btn.click()
                    print("[OK] 已确认清除所有筛选")
                    page.wait_for_load_state('networkidle')
                    time.sleep(2)

            # 展开分组
            print("展开分组...")
            page.click('text=一键展开')
            page.wait_for_load_state('networkidle')
            time.sleep(1)

            # 导出
            print("导出 Excel...")
            page.click('button:has-text("导出")')
            page.wait_for_selector('text=导出excel', timeout=10000)

            with page.expect_download(timeout=60000) as dl:
                page.click('text=导出excel')

            download = dl.value
            file_path = os.path.join(DOWNLOAD_DIR, download.suggested_filename)
            download.save_as(file_path)
            print(f"[OK] 下载完成: {file_path}")
            return file_path

        except Exception as e:
            print(f"[ERROR] 导出失败: {e}")
            return None
        finally:
            browser.close()
            print("浏览器已关闭")


# ============================================================
#  Phase 2 — 数据筛选
# ============================================================
def filter_data(source_file, plan_date, ver_date, ver_only=False):
    print("\n" + "=" * 60)
    print("Phase 2: 筛选数据")
    print("=" * 60)
    if ver_only:
        print(f"  条件: 版本排期 = {ver_date} (仅版本排期模式)")
    else:
        print(f"  条件 1: 计划生产排期 = {plan_date}")
        print(f"  条件 2: 版本排期     = {ver_date}")

    wb = openpyxl.load_workbook(source_file)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]

    col_ver = headers.index('版本排期') + 1
    col_plan = headers.index('计划生产排期') + 1 if '计划生产排期' in headers else None

    filtered_rows = []
    keep_indices = set()       # 0-based data row indices to keep
    original_count = 0

    for ri, row in enumerate(ws.iter_rows(min_row=1, values_only=True)):
        if ri == 0:
            filtered_rows.append(list(row))
        else:
            original_count += 1
            v_ver = str(row[col_ver - 1]) if row[col_ver - 1] else ''
            v_ver_clean = v_ver.replace('-', '')

            if ver_only:
                # 仅按版本排期筛选
                if ver_date in v_ver_clean:
                    filtered_rows.append(list(row))
                    keep_indices.add(ri - 1)
            else:
                v_plan = str(row[col_plan - 1]) if col_plan and row[col_plan - 1] else ''
                # 宽口径版本匹配：如果 ver_date 包含 plan_date，则允许匹配 20260703 或 20260703_upt
                ver_match = (ver_date in v_ver_clean) or (plan_date in ver_date and plan_date in v_ver_clean)
                if plan_date in v_plan.replace('-', '') and ver_match:
                    filtered_rows.append(list(row))
                    keep_indices.add(ri - 1)

    # 如果严格筛选无数据，且是单参数模式（即 plan_date == ver_date），则回退到仅按计划生产排期筛选
    if not ver_only and len(filtered_rows) <= 1 and plan_date == ver_date:
        print(f'[WARN] 未找到"版本排期"包含 {ver_date} 的数据。尝试仅按"计划生产排期 = {plan_date}"进行宽口径匹配...')
        filtered_rows = [filtered_rows[0]]  # 保留表头
        keep_indices.clear()
        for ri, row in enumerate(ws.iter_rows(min_row=1, values_only=True)):
            if ri > 0:
                v_plan = str(row[col_plan - 1]) if row[col_plan - 1] else ''
                if plan_date in v_plan.replace('-', ''):
                    filtered_rows.append(list(row))
                    keep_indices.add(ri - 1)

    wb.close()
    print(f"  原始: {original_count} 条 → 筛选后: {len(filtered_rows) - 1} 条")
    return filtered_rows, headers, original_count, keep_indices


# ============================================================
#  Phase 3 — 文件生成
# ============================================================
def save_main_file(source_file, keep_indices, version):
    """直接操作源文件副本：保留筛选行，删除其余 → 完整保留原始格式/样式/列宽"""
    print("\n" + "=" * 60)
    print("Phase 3: 文件生成")
    print("=" * 60)

    main_name = f"{BASE_VERSION}_{version}.xlsx"
    backup_name = f"{BASE_VERSION}_{version}_backup.xlsx"
    main_path = os.path.join(EXPORT_DIR, main_name)
    backup_path = os.path.join(EXPORT_DIR, backup_name)

    # 备份：上一版主文件
    if os.path.exists(main_path):
        shutil.copy2(main_path, backup_path)
        print(f"[OK] 已备份 → {backup_name}")
    else:
        # 如果是增量/热修复版本且主文件不存在，尝试寻找基线版本主文件进行备份
        if "_" in version:
            base_ver = version.split("_")[0]
            base_main_path = os.path.join(EXPORT_DIR, f"{BASE_VERSION}_{base_ver}.xlsx")
            if os.path.exists(base_main_path):
                shutil.copy2(base_main_path, backup_path)
                print(f"[OK] 未找到当前主文件，已自动复制基线版本主文件为备份 → {backup_name}")

    # 加载源文件，批量删除不符合条件的行（保留原始格式）
    wb = openpyxl.load_workbook(source_file)
    ws = wb.active

    # 快速合并连续行删除
    keep_rows = {idx + 2 for idx in keep_indices}
    r = ws.max_row
    while r >= 2:
        if r not in keep_rows:
            block_end = r
            block_start = r
            while block_start - 1 >= 2 and (block_start - 1) not in keep_rows:
                block_start -= 1
            ws.delete_rows(block_start, block_end - block_start + 1)
            r = block_start - 1
        else:
            r -= 1

    wb.save(main_path)
    wb.close()
    print(f"[OK] 主文件 → {main_name}  (保留 {len(keep_indices)} 条)")
    return main_path


def generate_upt_report(source_file, keep_indices, filtered_rows, col_tid, version):
    """增量报告：直接从源文件保留新增行 → 原始格式不丢失"""
    print("\n" + "=" * 60)
    print("增量变动报告")
    print("=" * 60)

    upt_name = f"{BASE_VERSION}_{version}_upt.xlsx"
    upt_path = os.path.join(EXPORT_DIR, upt_name)
    backup_path = os.path.join(EXPORT_DIR, f"{BASE_VERSION}_{version}_backup.xlsx")

    # 读取上期 Task
    prev_ids = set()
    if os.path.exists(backup_path):
        try:
            bw = openpyxl.load_workbook(backup_path, read_only=True)
            bws = bw.active
            bh = [c.value for c in bws[1]]
            bc = bh.index('Task编号')
            for row in bws.iter_rows(min_row=2, values_only=True):
                if row[bc] is not None:
                    prev_ids.add(str(row[bc]))
            bw.close()
            print(f"[INFO] 上期 Task: {len(prev_ids)} 条")
        except Exception as e:
            print(f"[WARN] 读取备份失败: {e}")

    # 建 Task编号 → 0-based data index 映射（从源文件）
    cur_ids = set()
    tid_to_source_idx = {}
    wb = openpyxl.load_workbook(source_file)
    ws = wb.active
    src_headers = [c.value for c in ws[1]]
    sc = src_headers.index('Task编号')
    for di in range(ws.max_row - 1):
        val = ws.cell(row=di + 2, column=sc + 1).value
        if val:
            tid_to_source_idx[str(val)] = di
    wb.close()

    # 确定新增行的 0-based data index
    new_data_indices = []
    for row in filtered_rows[1:]:
        tid = str(row[col_tid]) if col_tid < len(row) and row[col_tid] is not None else ''
        if tid:
            cur_ids.add(tid)
            if tid not in prev_ids and tid in tid_to_source_idx:
                new_data_indices.append(tid_to_source_idx[tid])

    print(f"[INFO] 本期 Task: {len(cur_ids)} 条, 新增: {len(new_data_indices)} 条")

    # 从源文件生成增量报告（保留原始格式）
    nw = openpyxl.load_workbook(source_file)
    nws = nw.active

    # 快速合并连续行删除
    keep_rows = {idx + 2 for idx in new_data_indices}
    r = nws.max_row
    while r >= 2:
        if r not in keep_rows:
            block_end = r
            block_start = r
            while block_start - 1 >= 2 and (block_start - 1) not in keep_rows:
                block_start -= 1
            nws.delete_rows(block_start, block_end - block_start + 1)
            r = block_start - 1
        else:
            r -= 1

    nw.save(upt_path)
    nw.close()
    print(f"[OK] 增量报告 → {upt_name}")
    return upt_path


# ============================================================
#  Phase 4 — 模板处理
# ============================================================
def _collect_simple(df, col, split_pattern=r'[;,；，、\s]+'):
    """通用列提取：按分隔符拆分 → 去重列表"""
    if col not in df.columns:
        return []
    items = []
    for text in df[col].dropna().astype(str):
        for part in re.split(split_pattern, text):
            p = part.strip()
            if p and p.lower() != 'nan':
                items.append(p)
    return list(dict.fromkeys(items))


def _collect_path(df, col, list_var, output_lines, account_suffix=""):
    """通用路径列处理：经办人 + Task编号 + 路径 → append 语句"""
    if col not in df.columns:
        return 0, 0
    cnt, skipped = 0, 0
    for _, row in df.iterrows():
        account = name_to_pinyin(row.get('经办人'))
        tid = str(row.get('Task编号', '')).strip() if not pd.isna(row.get('Task编号')) else ''
        parsed = parse_path_string(row.get(col))
        if not parsed:
            skipped += 1
            continue
        output_lines.append(f"{list_var}.append(['{account}{tid}{account_suffix}', {parsed}])")
        cnt += 1
    return cnt, skipped


def run_template_processing(main_path, version, output_name=None):
    print("\n" + "=" * 60)
    print("Phase 4: 模板处理")
    print("=" * 60)

    if output_name is None:
        output_name = OUTPUT_TXT

    if not os.path.exists(main_path):
        print(f"[ERROR] 文件不存在: {main_path}")
        sys.exit(1)

    try:
        df = pd.read_excel(main_path)
        print(f"[OK] 读取 {len(df)} 行, {len(df.columns)} 列")
    except Exception as e:
        print(f"[ERROR] 读取失败: {e}")
        sys.exit(1)

    out = []
    out.append(
        f"# 自动生成报告\n"
        f"# 版本: {BASE_VERSION}_{version}\n"
        f"# 源文件: {os.path.basename(main_path)}\n"
        f"# 模块: CLI, DLL, Table, HisRunTable, HisTable, Proc, ProcAdd, "
        f"TableDevelop, UpgradeRemark, Remark, Init\n\n"
    )

    # Task编号 汇总
    if 'Task编号' in df.columns:
        tid_to_info = {}
        for _, row in df.iterrows():
            tid = str(row.get('Task编号', '')).strip()
            operator = str(row.get('经办人', '')).strip()
            if pd.isna(row.get('经办人')) or operator.lower() == 'nan':
                operator = ''
            
            title = str(row.get('Task标题', '')).strip()
            if pd.isna(row.get('Task标题')) or title.lower() == 'nan':
                title = ''
                
            if tid and tid.lower() != 'nan':
                tid_to_info[tid] = (operator, title)
        
        tids = sorted(list(tid_to_info.keys()))
        if tids:
            out.append(f"# Task编号 (共 {len(tids)} 项):\n")
            for t in tids:
                out.append(f"{t}\n")
            out.append("\n")
            
            out.append("# Task编号详细\n")
            for t in tids:
                op, title = tid_to_info[t]
                if title and op:
                    out.append(f"{t} {title}  {op}\n")
                elif title:
                    out.append(f"{t} {title}\n")
                elif op:
                    out.append(f"{t}  {op}\n")
                else:
                    out.append(f"{t}\n")
            out.append("\n")

    tasks = [
        ("CLI",              lambda: _do_simple(df, "cli", "cli_list", out)),
        ("DLL",              lambda: _do_dll(df, out)),
        ("Table",            lambda: _do_table(df, out)),
        ("HisRunTable",      lambda: _do_path(df, "历史run表名/表文件", "table_his_list_run", out)),
        ("HisTable",         lambda: _do_path(df, "历史his表名/表文件", "table_his_list_his", out)),
        ("Proc",             lambda: _do_path(df, "ProcUpd", "proc_list", out)),
        ("ProcAdd",          lambda: _do_path(df, "ProcAdd", "proc_list", out, account_suffix="_new")),
        ("TableDevelop",     lambda: _do_simple_plain(df, "table_develop下文件", out)),
        ("UpgradeRemark",    lambda: _do_plain_with_tid(df, "升级备注", out)),
        ("Remark",           lambda: _do_plain_with_tid(df, "备注", out)),
        ("Init",             lambda: _do_simple_plain(df, "init", out)),
    ]

    for name, func in tasks:
        print(f"  [{name}] ...", end=" ")
        func()
        print("OK")

    # 写输出
    if out:
        txt_path = os.path.join(EXPORT_DIR, output_name)
        with open(txt_path, 'w', encoding='utf-8') as f:
            f.writelines(out)
        print(f"\n[OK] 模板输出 → {output_name}")
    else:
        print("\n[WARN] 所有任务均未产生输出")


def _do_simple(df, col, list_name, out):
    """简单列表（去重 + 带变量名）"""
    items = _collect_simple(df, col, r'[;,；，、\s]+')
    if not items:
        return
    out.append(f"# === {col} (共 {len(items)} 项) ===\n")
    out.append(f"{list_name} = {items}\n\n")
    out.append("\n".join(items) + "\n\n")


def _do_dll(df, out):
    """DLL 特殊处理：去掉 .dll 后缀"""
    raw = df.get("lbm")
    if raw is None:
        return
    items = []
    for text in raw.dropna().astype(str):
        for part in re.split(r'[;,；，、\s]+', text):
            p = part.replace('.dll', '').strip()
            if p:
                items.append(p)
    items = list(dict.fromkeys(items))
    if not items:
        return
    out.append(f"# === DLL (共 {len(items)} 项) ===\n")
    out.append(f"lbm_list = {items}\n\n")
    out.append("\n".join(items) + "\n\n")


def _do_path(df, col, list_var, out, account_suffix=""):
    """路径列"""
    lines = []
    cnt, skipped = _collect_path(df, col, list_var, lines, account_suffix)
    if lines:
        out.append(f"# === {col} (有效 {cnt} 条, 跳过 {skipped} 条) ===\n")
        out.append("\n".join(lines) + "\n\n")


def _do_table(df, out):
    """Table 列特殊处理：同时检查 table名/table文件名/备份表名 和 table 列"""
    lines = []
    col1 = "table名/table文件名/备份表名"
    col2 = "table"
    
    cnt1, skipped1 = 0, 0
    if col1 in df.columns:
        cnt1, skipped1 = _collect_path(df, col1, "table_run_list", lines)
        
    cnt2, skipped2 = 0, 0
    if col2 in df.columns:
        cnt2, skipped2 = _collect_path(df, col2, "table_run_list", lines)
        
    cnt = cnt1 + cnt2
    skipped = skipped1 + skipped2
    if lines:
        out.append(f"# === Table (有效 {cnt} 条, 跳过 {skipped} 条) ===\n")
        out.append("\n".join(lines) + "\n\n")


def _do_simple_plain(df, col, out):
    """纯文本列表（无变量名）"""
    items = _collect_simple(df, col)
    if not items:
        return
    out.append(f"# === {col} (共 {len(items)} 项) ===\n")
    out.append("\n".join(items) + "\n\n")


def _do_plain_with_tid(df, col, out):
    """带 Task 编号前缀的文本列表"""
    if col not in df.columns:
        return
    lines = []
    for _, row in df.iterrows():
        val = row.get(col)
        if pd.isna(val) or not str(val).strip() or str(val).lower() == 'nan':
            continue
        tid = str(row.get('Task编号', '')).strip() if not pd.isna(row.get('Task编号')) else ''
        text = str(val).strip()
        if tid:
            lines.append(f"{tid}      {text}")
        else:
            lines.append(text)
            
    if lines:
        out.append(f"# === {col} (共 {len(lines)} 项) ===\n")
        out.append("\n".join(lines) + "\n\n")


# ============================================================
#  汇总
# ============================================================
def print_summary(version, plan_date, ver_date, main_path, upt_path,
                  filtered_count, original_count):
    print("\n" + "=" * 60)
    print("任务完成 — 汇总")
    print("=" * 60)
    print(f"  版本号:         {version}")
    if plan_date:
        print(f"  计划生产排期:   {plan_date}")
    else:
        print(f"  计划生产排期:   (仅版本排期模式，未筛选)")
    print(f"  版本排期:       {ver_date}")
    print(f"  生成时间:       {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"  数据:           {original_count} → {filtered_count} 条")
    print(f"  主文件:         {os.path.basename(main_path)}")
    print(f"  增量报告:       {os.path.basename(upt_path)}")
    print(f"  模板输出(全量): {OUTPUT_TXT}")
    print(f"  模板输出(增量): {OUTPUT_UPT_TXT}")
    print(f"  输出目录:       {EXPORT_DIR}")
    print("=" * 60)


def clean_old_versions(export_dir, current_version):
    """自动清理目录下非当前版本日期的文件和文件夹"""
    print("\n" + "=" * 60)
    print("清理历史版本文件")
    print("=" * 60)
    pattern_prefix = BASE_VERSION.lower()
    base_date = current_version[:8]
    cleaned_count = 0
    for filename in os.listdir(export_dir):
        fn_lower = filename.lower()
        if fn_lower.startswith(pattern_prefix + "_") or fn_lower.startswith(pattern_prefix + "-"):
            if base_date.lower() not in fn_lower:
                file_path = os.path.join(export_dir, filename)
                try:
                    if os.path.isdir(file_path):
                        shutil.rmtree(file_path)
                        print(f"[CLEAN] 已删除历史版本文件夹: {filename}")
                    else:
                        os.remove(file_path)
                        print(f"[CLEAN] 已删除历史版本文件: {filename}")
                    cleaned_count += 1
                except Exception as e:
                    print(f"[WARN] 无法删除 {filename}: {e}")
    if cleaned_count == 0:
        print("  没有发现需要清理的历史版本文件")
    else:
        print(f"  已清理 {cleaned_count} 个历史版本文件/文件夹")


# ============================================================
#  主流程
# ============================================================
def main():
    version, plan_date, ver_date, ver_only = parse_version()

    print("\n" + "=" * 60)
    print(f"国泰海通 SPB 自动化导出工具")
    if ver_only:
        print(f"版本基线: {BASE_VERSION}  |  版本排期: {ver_date} (仅版本排期模式)")
    else:
        print(f"版本基线: {BASE_VERSION}  |  计划生产排期: {plan_date}  |  版本排期: {ver_date}")
    print("=" * 60)

    ensure_dir(EXPORT_DIR)

    # Phase 1: 抓取
    source_file = export_from_platform(version)
    if not source_file:
        print("[ERROR] Phase 1 失败，退出")
        sys.exit(1)

    # Phase 2: 筛选
    filtered_rows, headers, original_count, keep_indices = filter_data(source_file, plan_date, ver_date, ver_only)
    if len(filtered_rows) <= 1:
        print("[WARN] 无符合条件数据，退出")
        sys.exit(0)

    col_tid = headers.index('Task编号')

    # Phase 3: 存文件（直接操作源文件保留格式）
    main_path = save_main_file(source_file, keep_indices, version)
    upt_path = generate_upt_report(source_file, keep_indices, filtered_rows, col_tid, version)

    # Phase 4: 模板处理（全量 + 增量）
    print("\n  [全量模式] 处理主文件 -> excel_output.txt")
    run_template_processing(main_path, version)
    if os.path.exists(upt_path):
        try:
            upt_wb = openpyxl.load_workbook(upt_path, read_only=True)
            upt_ws = upt_wb.active
            upt_row_count = upt_ws.max_row - 1 if upt_ws.max_row else 0
            upt_wb.close()
            if upt_row_count > 0:
                print("\n  [增量模式] 处理增量报告 -> upt_excel_output.txt")
                run_template_processing(upt_path, version, OUTPUT_UPT_TXT)
            else:
                print(f"\n  [增量模式] 跳过：增量报告无新增数据 ({upt_row_count} 条)")
        except Exception as e:
            print(f"\n  [增量模式] 跳过：读取增量报告失败 ({e})")
    else:
        print(f"\n  [增量模式] 跳过：增量报告不存在")

    # 汇总
    print_summary(version, plan_date, ver_date, main_path, upt_path,
                  len(filtered_rows) - 1, original_count)

    # 自动清理历史版本
    clean_old_versions(EXPORT_DIR, version)


if __name__ == "__main__":
    main()
