import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
import os
import json
import sys
import time
import re

# Patch openpyxl to ignore 'id' attribute in DataValidation (WPS compatibility fix)
original_init = DataValidation.__init__
def patched_init(self, *args, **kwargs):
    kwargs.pop('id', None)
    original_init(self, *args, **kwargs)
DataValidation.__init__ = patched_init

def normalize_system_name(name):
    if not name:
        return ""
    return name.strip().replace("系统", "")

def main():
    cache_dir = "/Users/wujin/.gemini/antigravity/state/bcp_tracker"
    xlsx_path = os.path.join(cache_dir, "downloaded_workbook.xlsx")
    
    if not os.path.exists(xlsx_path):
        print(f"Error: Workbook not found at {xlsx_path}")
        sys.exit(1)
        
    print("Loading workbook...")
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    
    # 1. Parse BCP sheet for downstream systems status
    sheet = wb['BCP历史迁移']
    rows = list(sheet.iter_rows(values_only=True))
    
    systems = {
        "大数据平台": 14,
        "营运管理平台": 19,
        "非现场监控": 25,
        "信用业务管理系统": 31,
        "反洗钱监控系统": 37,
        "正回购业务风险监控系统": 43
    }
    
    # 2. Parse '其他' sheet for system contacts and BCP demands
    contacts = {}
    demands = []
    if '其他' in wb.sheetnames:
        other_sheet = wb['其他']
        for row in other_sheet.iter_rows(values_only=True):
            if len(row) >= 4:
                req_id = str(row[1]).strip() if row[1] else ''
                sys_or_desc = str(row[2]).strip() if row[2] else ''
                contact_or_remark = str(row[3]).strip() if row[3] else ''
                
                if re.match(r'^R\d+$', req_id):
                    demands.append({
                        'id': req_id,
                        'desc': sys_or_desc,
                        'remark': contact_or_remark
                    })
                elif sys_or_desc and contact_or_remark and sys_or_desc != '系统' and contact_or_remark != '负责人' and sys_or_desc != '需求描述':
                    contacts[sys_or_desc] = contact_or_remark
                    
    # Map BCP systems to contacts using normalized names
    system_contacts = {}
    for bcp_sys_name in systems.keys():
        norm_bcp = normalize_system_name(bcp_sys_name)
        matched_contact = "未登记"
        for other_sys_name, contact_val in contacts.items():
            norm_sheet3 = normalize_system_name(other_sys_name)
            if norm_bcp == norm_sheet3 or norm_bcp in norm_sheet3 or norm_sheet3 in norm_bcp:
                matched_contact = contact_val
                break
        system_contacts[bcp_sys_name] = matched_contact
        
    results = {}
    for sys_name, start_col in systems.items():
        results[sys_name] = []
        is_big_data = (sys_name == "大数据平台")
        
        use_idx = start_col + 1
        plan_idx = start_col + 2
        prog_idx = start_col + 3
        rem_idx = start_col + 4
        use_case_idx = None if is_big_data else start_col + 5
        timing_idx = None if is_big_data else start_col
        
        for row in rows[3:]:
            table_name = row[3]
            table_desc = row[5]
            change_stat = row[4]
            owner = row[8]
            
            if not table_name:
                continue
                
            in_use = row[use_idx]
            if in_use == '是':
                prog = row[prog_idx] or "未开始"
                plan = row[plan_idx] or "否"
                remark = row[rem_idx] or ""
                use_case = row[use_case_idx] if use_case_idx else ""
                timing = row[timing_idx] if timing_idx else ""
                
                results[sys_name].append({
                    "table_name": table_name,
                    "table_desc": table_desc,
                    "change_stat": change_stat,
                    "owner": owner,
                    "plan_migrate": plan,
                    "progress": prog,
                    "remark": remark,
                    "use_case": use_case,
                    "timing": timing
                })
                
    # 3. Save JSON Report
    with open(os.path.join(cache_dir, "systems_status.json"), "w", encoding="utf-8") as f:
        json.dump({"systems": results, "contacts": system_contacts, "demands": demands}, f, ensure_ascii=False, indent=2)
        
    # 4. Generate system list for reminders
    sys_reminders = []
    sys_order = ['信用业务管理系统', '非现场监控', '营运管理平台', '大数据平台', '反洗钱监控系统', '正回购业务风险监控系统']
    involved_people = set()
    
    for sname in sys_order:
        tables = results[sname]
        total_in_use = len(tables)
        
        # Filter tables that are planned to migrate to count statuses
        plan_tables = [t for t in tables if t['plan_migrate'] == '是']
        plan_migrate_count = len(plan_tables)
        
        # Calculate status counts for planned tables
        status_counts = {}
        for t in plan_tables:
            status_counts[t['progress']] = status_counts.get(t['progress'], 0) + 1
            
        # Format status counts in a clean string
        if plan_migrate_count > 0:
            status_parts = []
            status_order = ['已上线', '测试完', '测试中', '开发中', '未开始']
            other_statuses = [k for k in status_counts.keys() if k not in status_order]
            full_order = status_order + other_statuses
            
            for status in full_order:
                if status in status_counts:
                    status_parts.append(f"{status} {status_counts[status]}个")
            prog_str = "目前进度 " + "、".join(status_parts)
        else:
            prog_str = "目前无迁移计划"
            
        contact_str = system_contacts.get(sname, '')
        teachers = []
        at_tags = []
        if contact_str and contact_str != '未登记':
            names = contact_str.replace('、', ' ').replace(',', ' ').split()
            for name in names:
                last_name = name[0]
                teachers.append(f'{last_name}老师')
                at_tags.append(f'@{name}')
                if len(name) >= 2:
                    involved_people.add(name)
            teachers_phrase = '、'.join(teachers)
            at_phrase = ' '.join(at_tags)
        else:
            teachers_phrase = '相关老师'
            at_phrase = ''
            
        sys_reminders.append({
            "sys_name": sname,
            "total_in_use": total_in_use,
            "plan_migrate": plan_migrate_count,
            "progress": prog_str,
            "teachers": teachers_phrase,
            "at": at_phrase
        })

    # Get formatted system sentences for Version A (only keep Version A)
    all_sentences = []
    for idx, r in enumerate(sys_reminders, start=1):
        sentence = f"{idx}. {r['sys_name']}：共涉及{r['total_in_use']}张在用表，计划迁移{r['plan_migrate']}张表，{r['progress']}，请{r['teachers']}确认近期的适配及测试验证安排；{r['at']}".strip()
        all_sentences.append(sentence)
        
    sec_all_systems_str = "\n  ".join(all_sentences)

    # Parse and compile migration demand items from '其他' sheet
    demand_lines = []
    for d in demands:
        target_sys = ""
        if "清算" in d['desc']:
            target_sys = "集中清算"
        elif "存管" in d['desc']:
            target_sys = "三方存管"
        elif "参数" in d['desc']:
            target_sys = "参数中心"
            
        involved_names = []
        source_names = contacts.get("集中交易", "").replace("、", " ").split()
        for name in source_names:
            if len(name) >= 2:
                involved_names.append(name)
                
        if target_sys:
            target_names = contacts.get(target_sys, "").replace("、", " ").split()
            for name in target_names:
                if len(name) >= 2:
                    involved_names.append(name)
                    
        # Deduplicate
        unique_names = []
        for n in involved_names:
            if n not in unique_names:
                unique_names.append(n)
                
        teachers = []
        at_tags = []
        for name in unique_names:
            last_name = name[0]
            teachers.append(f"{last_name}老师")
            at_tags.append(f"@{name}")
            involved_people.add(name)
            
        teachers_phrase = "、".join(teachers)
        at_phrase = " ".join(at_tags)
        
        line = f"{d['desc']}（{d['id']}）：状态为“{d['remark']}”，请{teachers_phrase}关注；{at_phrase}"
        demand_lines.append((d['remark'], line))

    active_demands = []
    ended_demands = []
    for status, line in demand_lines:
        if status in ["结束", "终止"]:
            ended_demands.append(line)
        else:
            active_demands.append(line)
            
    sec_demands_parts = []
    if active_demands:
        sec_demands_parts.append("（一） 进行中需求：\n  " + "\n  ".join(f"{idx}. {line}" for idx, line in enumerate(active_demands, start=1)))
    if ended_demands:
        sec_demands_parts.append("（二） 已上线/已结束需求：\n  " + "\n  ".join(f"{idx}. {line}" for idx, line in enumerate(ended_demands, start=1)))
        
    sec_demands_str = "\n\n  ".join(sec_demands_parts)
    
    # Resolve all recipients
    static_recipients = [
        "黄志昌", "杨蓉渊", "柴恒", "李骎", "周静", "胡玲杰", 
        "揭由翔", "丁伟", "高健", "李萍萍", "曾娟", "茆莹莹", "刘青"
    ]
    all_recipients = sorted(list(involved_people.union(static_recipients)))
    recipients_str = " ".join(all_recipients)
    
    cc_list = ["姜婷婷", "钱维佳", "吴保杰", "彭伟", "纪飞", "赵永杰", "周哲博", "周尤珠"]
    cc_str = " ".join(cc_list)

    date_str = time.strftime("%Y%m%d")
    
    # Draft Mail Contents (Only keeping Version A)
    mail_draft = f"""# {date_str}-进度-集中交易历史数据BCP文件迁移项目下游系统适配进度跟踪

## 1. 邮件起草草稿 (版本 A：正式严谨版)

Subject (邮件主题)：
【进度提醒】关于集中交易历史数据BCP文件迁移项目下游系统适配进度跟踪-{date_str}

To (收件人)：
{recipients_str}

CC (抄送)：
{cc_str}

Body (正文)：
各位老师：

  针对集中交易历史数据BCP文件迁移专项项目的整体进度，目前项目已进入下游系统适配及联调的关键阶段。为了确保整体项目能够按计划稳步推进，现将各下游系统在用表迁移及开发的最新进度梳理如下，请各位老师对照并予以关注：

  一、 下游系统适配进度：
  {sec_all_systems_str}

  为了保障下游系统能按期适配上线，避免数据断档，请各位老师在在线文档中及时更新最新进展，并对未开始/开发中的表确认排期。该项目须确保于第三季度内完成交付，如有疑问或需要上游协助，请随时沟通。

在线文档如下：
https://www.gtht.com.cn/jswps/weboffice/l/s10TcLAYXLAeL

---

## 二、 下游系统表级明细
"""

    for sys_name, tables in results.items():
        contact_person = system_contacts.get(sys_name, "未登记")
        mail_draft += f"\n### {sys_name} (负责人：{contact_person} | {len(tables)} 张在用表)\n\n"
        if not tables:
            mail_draft += "无在用表或未登记。\n"
            continue
            
        mail_draft += "| 序号 | 表名 | 表说明 | 拟变更统计 | 计划迁移 | 进度 | 备注 |\n"
        mail_draft += "| --- | --- | --- | --- | --- | --- | --- |\n"
        visible_idx = 0
        for t in tables:
            if t['plan_migrate'] == '否':
                continue
            visible_idx += 1
            display_progress = f"**{t['progress']}**"
            mail_draft += f"| {visible_idx} | `{t['table_name']}` | {t['table_desc'] or ''} | {t['change_stat'] or ''} | {t['plan_migrate']} | {display_progress} | {t['remark']} |\n"
        mail_draft += "\n"
        
    # Append sign-off to mail_draft
    mail_draft += "顺祝商祺！\n"

    # Write to local cache
    md_report_path = os.path.join(cache_dir, "systems_status.md")
    with open(md_report_path, "w", encoding="utf-8") as f:
        f.write(mail_draft)

    # Write raw subject and body to separate files for automated drafting
    raw_subject = f"【进度提醒】关于集中交易历史数据BCP文件迁移项目下游系统适配进度跟踪-{date_str}"
    raw_body = f"""各位老师：
 
  针对集中交易历史数据BCP文件迁移专项项目的整体进度，目前项目已进入下游系统适配及联调的关键阶段。为了确保整体项目能够按计划稳步推进，现将各下游系统在用表迁移及开发的最新进度梳理如下，请各位老师对照并予以关注：
 
  一、 下游系统适配进度：
  {sec_all_systems_str}
 
  为了保障下游系统能按期适配上线，避免数据断档，请各位老师在在线文档中及时更新最新进展，并对未开始/开发中的表确认排期。该项目须确保于第三季度内完成交付，如有疑问或需要上游协助，请随时沟通。
 
在线文档如下：
https://www.gtht.com.cn/jswps/weboffice/l/s10TcLAYXLAeL
 
顺祝商祺！"""
    
    with open(os.path.join(cache_dir, "mail_subject.txt"), "w", encoding="utf-8") as f:
        f.write(raw_subject)
    with open(os.path.join(cache_dir, "mail_body.txt"), "w", encoding="utf-8") as f:
        f.write(raw_body)
        
    # Write to User's Obsidian Project Directory with the new naming scheme
    obsidian_dir = "/Volumes/Macintosh HD_Data/obsidian/100_Projects/进度跟踪"
    os.makedirs(obsidian_dir, exist_ok=True)
    obsidian_file_path = os.path.join(obsidian_dir, f"{date_str}-进度-集中交易历史数据BCP文件迁移项目下游系统适配进度跟踪.md")
    with open(obsidian_file_path, "w", encoding="utf-8") as f:
        f.write(mail_draft)
        
    print(f"✅ Excel workbook parsed successfully!")
    print(f"Obsidian report saved to: {obsidian_file_path}")

if __name__ == "__main__":
    main()
